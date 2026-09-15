# -*- coding: utf-8 -*-
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Malgun Gothic'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Malgun Gothic'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Blue 900
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Malgun Gothic'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r = p.add_run(text)
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Malgun Gothic'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r = p.add_run(text)
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_callout(doc, text, title="핵심 요약"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"💡 {title}: ")
    r_title.font.name = 'Malgun Gothic'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

    r_body = p.add_run(text)
    r_body.font.name = 'Malgun Gothic'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def main():
    doc = docx.Document()

    # 페이지 여백 설정
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. 문서 헤더 / 표지 블록
    tbl_header = doc.add_table(rows=1, cols=1)
    tbl_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_h = tbl_header.cell(0, 0)
    cell_h.width = Inches(6.8)
    set_cell_background(cell_h, "0F172A") # Slate 900
    set_cell_margins(cell_h, top=240, bottom=240, left=240, right=240)

    p_tag = cell_h.paragraphs[0]
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("BEAUSCREATORS PRODUCT SPECIFICATION REPORT")
    r_tag.font.name = 'Malgun Gothic'
    r_tag.font.size = Pt(9)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8) # Sky 400

    p_t = cell_h.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(6)
    r_t = p_t.add_run("RunNow 1:1 전담 AI 케어팀\n4대 에이전트 상세 프로필 및 역할 명세서")
    r_t.font.name = 'Malgun Gothic'
    r_t.font.size = Pt(17)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_sub = cell_h.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    r_sub = p_sub.add_run("운동 PT(남/여) · 식단 영양사 · 메디컬 닥터 3대 전문 영역 | 음성 TTS 스펙 | 멀티 에이전트 티키타카 시나리오")
    r_sub.font.name = 'Malgun Gothic'
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(8)
    p_meta.paragraph_format.space_after = Pt(14)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_m = p_meta.add_run("보고 대상: 이건우 대표님 (BeausCreators CEO) | 작성: BSC 비서실 & RunNow 개발 셀 | 일자: 2026.09.14")
    r_m.font.name = 'Malgun Gothic'
    r_m.font.size = Pt(8.5)
    r_m.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 1. 기획 개요 및 포지셔닝
    add_heading_1(doc, "1. 기획 개요 및 케어팀 도입 목적")
    add_body_p(doc, "본 문서는 RunNow(나이키 러닝클럽 x 다마고치 x 온디바이스 AI 러닝 OS)의 인앱 프리미엄 가치를 극대화하기 위해 탑재되는 1:1 전담 AI 케어팀의 4대 코치 프로필, 음성 오디오 규격, 데이터 연동 인터랙션 및 상호 협업 파이프라인을 정의합니다.")
    add_callout(doc, "오프라인 1:1 PT는 회당 6~8만 원에 달하지만, RunNow는 '남/여 전담 트레이너 + 식약처 영양사 + 스포츠 재활 닥터' 3개 전문 영역의 4대 페르소나를 스마트폰 속에 원스톱 제공하여 월 50만 원 상당의 종합 피트니스 클리닉 경험을 혁신합니다.", "핵심 가치 제안")

    # 2. 4대 전담 에이전트 총괄 라인업 요약표
    add_heading_1(doc, "2. 4대 전담 에이전트 총괄 라인업 요약표")
    
    tbl_summary = doc.add_table(rows=5, cols=5)
    tbl_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_summary)

    headers = ["영역", "에이전트명 (ID)", "페르소나 성향", "핵심 담당 역할", "연동 핵심 데이터"]
    for i, h in enumerate(headers):
        cell = tbl_summary.cell(0, i)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = 'Malgun Gothic'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    summary_rows = [
        ("운동 PT (남)", "코치 레오\n(leo)", "파이팅 넘치고 에너제틱한 남성 트레이너", "하체/코어 한계돌파, 강한 동기부여, 고강도 인터벌 러닝", "카메라 모션인식(자세/횟수),\nGPS 실측 거리/페이스"),
        ("운동 PT (여)", "코치 루나\n(luna)", "섬세하고 스마트한 여성 페이스메이커", "지속 가능한 러닝 페이스 조절, 유연성 및 섬세한 자세 교정", "하버사인 정밀 GPS 속도,\n카메라 33개 관절 랜드마크"),
        ("식단 영양", "영양 코치 엘리\n(ellie)", "공감형 밸런스 케어 전문 영양사", "일일 식단 체크리스트, 칼로리/탄단지 계산, 치팅 만회 가이드", "식약처 공공데이터 5만+ DB,\n운동 소모 칼로리(METs)"),
        ("메디컬 닥터", "닥터 케이\n(drkay)", "신뢰감 있고 침착한 스포츠 재활 전문의", "운동 전/후 컨디션 스캔, 관절/근육 통증 감지, 과훈련 방지", "질병관리청 건강지표,\n유저 피로도/통증 트래커")
    ]

    for row_idx, data in enumerate(summary_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = tbl_summary.cell(row_idx, col_idx)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Malgun Gothic'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    add_body_p(doc, "", space_after=6)

    # 3. 에이전트별 세부 프로필 및 인터랙션 스펙
    add_heading_1(doc, "3. 에이전트별 세부 프로필 및 인터랙션 스펙")

    # 3.1 코치 레오
    add_heading_2(doc, "3.1 코치 레오 (Leo) — 남성 전담 파워 & 인터벌 PT")
    add_bullet_p(doc, " '대표님의 한계를 깨뜨리고 폭발적인 성취감을 만드는 파워 러닝 코치'", "슬로건:")
    add_bullet_p(doc, " 남성 트레이너 보이스 (Pitch: 0.9, Rate: 1.05) — 당당하고 활력 넘치며 신뢰감을 주는 저음 톤", "음성 TTS 스펙:")
    add_bullet_p(doc, " 유저와 자연어 대화를 통해 주간 운동 목표(러닝 거리, 소모 칼로리) 자동 수립 및 캘린더 등록", "핵심 기능 1:")
    add_bullet_p(doc, " 러닝 1km 구간마다 실시간 음성 페이스 푸시 ('대표님, 페이스 완벽합니다! 300미터만 더 피치 올립시다!')", "핵심 기능 2:")
    add_bullet_p(doc, " 카메라 모션인식 기반 스쿼트/푸시업 실시간 카운트 및 '하나! 둘! 마지막 하나 더!' 파이팅 구호", "핵심 기능 3:")
    add_bullet_p(doc, " '대표님, 오늘 점심 삼겹살 든든하게 드신 만큼 힘 넘치시죠? 오늘 저녁 러닝 15분만 더 태우고 깔끔하게 퉁치시죠! 준비되셨습니까?'", "대표 발화:")

    # 3.2 코치 루나
    add_heading_2(doc, "3.2 코치 루나 (Luna) — 여성 전담 디테일 & 페이스 PT")
    add_bullet_p(doc, " '지치지 않고 꾸준히, 건강하고 즐거운 러닝 라이프를 만드는 페이스메이커'", "슬로건:")
    add_bullet_p(doc, " 여성 코치 보이스 (Pitch: 1.15, Rate: 0.98) — 맑고 차분하며 세심하게 감싸주는 친절한 톤", "음성 TTS 스펙:")
    add_bullet_p(doc, " 초보 러너 및 체력 부담을 느끼는 유저를 위한 저강도 존2(Zone 2) 유산소 스케줄 설계", "핵심 기능 1:")
    add_bullet_p(doc, " 러닝 호흡법, 발뒤꿈치 착지(힐스트라이크) 방지 및 무릎 충격 완화 자세 디테일 피드백", "핵심 기능 2:")
    add_bullet_p(doc, " 당일 유저 컨디션이 저하되었을 때 무리한 러닝 대신 가벼운 인터벌 파워워킹으로 자동 전환", "핵심 기능 3:")
    add_bullet_p(doc, " '대표님, 오늘 날씨가 조금 무덥네요. 무리해서 빨리 달리기보다는 호흡에 집중하면서 일정한 템포를 지켜봐요. 제가 옆에서 1km마다 계속 페이스 체크해 드릴게요!'", "대표 발화:")

    # 3.3 영양 코치 엘리
    add_heading_2(doc, "3.3 영양 코치 엘리 (Ellie) — 식약처 공인 식단 & 뉴트리션 코치")
    add_bullet_p(doc, " '자책감 없는(No-Guilt) 현실 식단 관리로 평생 지속 가능한 영양 밸런스 완성'", "슬로건:")
    add_bullet_p(doc, " 여성 영양사 보이스 (Pitch: 1.1, Rate: 1.0) — 따뜻하고 상냥하며 과학적 근거를 명쾌하게 짚어주는 톤", "음성 TTS 스펙:")
    add_bullet_p(doc, " 식약처 통합영양성분 5만+ DB 연동: 유저가 음식명 입력 시 1인분당 열량/탄단지/나트륨 0.1초 자동 산출", "핵심 기능 1:")
    add_bullet_p(doc, " 일상 실천형 체크리스트 관리: 기상 미온수 300ml, 밥 2/3공기 덜기, 국물 남기기, 야식 방어 도장", "핵심 기능 2:")
    add_bullet_p(doc, " 치팅/과식 솔직 고백 시 즉각 심리 안도감 제공 및 저녁 식단 자동 다운사이징(Rescue) 실행", "핵심 기능 3:")
    add_bullet_p(doc, " '대표님, 점심 회식 맛있게 드셨나요? 괜찮아요! 맛있게 먹은 한 끼는 절대 살로 가지 않아요. 대신 저녁엔 칼륨이 풍부한 샐러드와 달걀 2개로 가볍게 밸런스를 맞춰봐요. 제가 체크리스트 조정해 둘게요!'", "대표 발화:")

    # 3.4 닥터 케이
    add_heading_2(doc, "3.4 닥터 케이 (Dr. Kay) — 메디컬 & 컨디션 세이프티 닥터")
    add_bullet_p(doc, " '부상 없는 운동이 최고의 운동입니다. 과학적 의학 데이터로 대표님의 몸을 안전하게 보호합니다.'", "슬로건:")
    add_bullet_p(doc, " 남성 전문의 보이스 (Pitch: 0.85, Rate: 0.95) — 침착하고 지적이며 단단한 신뢰감을 주는 메디컬 톤", "음성 TTS 스펙:")
    add_bullet_p(doc, " 운동 전/후 유저 컨디션 스캔: 수면 시간, 발목/무릎 관절 뻐근함, 근육통 유무 자동 체크", "핵심 기능 1:")
    add_bullet_p(doc, " 통증 감지 시 부상 방지 알림 발송 및 고강도 운동을 폼롤러 마사지/동적 스트레칭으로 자동 치환", "핵심 기능 2:")
    add_bullet_p(doc, " 과훈련(Overtraining) 세이프티 가드: 주간 러닝 마일리지가 급격히 상승할 경우 강제 휴식일(Rest Day) 권고", "핵심 기능 3:")
    add_bullet_p(doc, " '안녕하십니까 대표님, 닥터 케이입니다. 지난 이틀 연속 러닝으로 무릎 연골에 피로가 감지됩니다. 오늘은 뛰지 마시고 따뜻한 족욕과 종아리 스트레칭 10분만 진행하십시오. 안전이 최우선입니다.'", "대표 발화:")

    # 4. 멀티 에이전트 실시간 티키타카(Rescue) 시나리오
    add_heading_1(doc, "4. 멀티 에이전트 실시간 티키타카(Rescue) 시나리오")
    add_body_p(doc, "유저가 식단이나 운동 데이터를 입력했을 때, 3개 영역의 에이전트들이 상호 데이터를 참조하며 실시간으로 대화형 케어를 제공하는 실제 시나리오입니다:")

    tbl_talk = doc.add_table(rows=1, cols=1)
    tbl_talk.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_talk = tbl_talk.cell(0, 0)
    cell_talk.width = Inches(6.8)
    set_cell_background(cell_talk, "F8FAFC")
    set_cell_margins(cell_talk, top=160, bottom=160, left=180, right=180)

    p_sc1 = cell_talk.paragraphs[0]
    p_sc1.paragraph_format.space_before = Pt(2)
    p_sc1.paragraph_format.space_after = Pt(4)
    r1 = p_sc1.add_run("📍 [실제 대화 흐름: 유저가 점심에 피자 3조각을 먹었다고 고백한 경우]\n\n")
    r1.font.name = 'Malgun Gothic'
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    dialogues = [
        ("🥗 영양사 엘리", "대표님, 피자 정말 맛있게 드셨나요? 죄책감 갖지 마세요! 오늘 저녁은 탄수화물을 덜어내고 단백질 위주 가벼운 샐러드로 밸런스를 맞춰드릴게요. (저녁 식단 자동 다운사이징)", "0284C7"),
        ("🏋️ 코치 레오", "오, 엘리 쌤 말대로 점심 든든히 드셨으니 오늘 에너지 풀 충전되셨겠네요! 그럼 오늘 퇴근길 저녁 러닝에 15분 인터벌 코스 제가 얹어두겠습니다. 저랑 땀 쫙 빼시죠! (운동 스케줄 자동 연장)", "16A34A"),
        ("🩺 닥터 케이", "레오 코치 플랜 좋습니다. 다만 대표님 어제 무릎 피로도가 60%로 기록되었으니, 러닝 전 발목 회전과 달리기 후 아이싱 쿨다운 5분을 필수 루틴으로 추가하겠습니다. (부상 방지 가드)", "D97706")
    ]

    for speaker, content, color_hex in dialogues:
        p_d = cell_talk.add_paragraph()
        p_d.paragraph_format.space_before = Pt(2)
        p_d.paragraph_format.space_after = Pt(4)
        r_spk = p_d.add_run(f"• {speaker}: ")
        r_spk.font.name = 'Malgun Gothic'
        r_spk.font.size = Pt(9)
        r_spk.font.bold = True
        r_spk.font.color.rgb = RGBColor(int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16))

        r_cnt = p_d.add_run(content)
        r_cnt.font.name = 'Malgun Gothic'
        r_cnt.font.size = Pt(9)
        r_cnt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 5. UI/UX 아바타 에셋 디자인 및 파일 저장 가이드
    add_heading_1(doc, "5. UI/UX 아바타 에셋 디자인 및 파일 저장 가이드")
    add_body_p(doc, "본사 전사 에이전트 자산과 분리하여, RunNow 앱 단독 배포(PWA/웹) 시에도 자산이 온전히 유지되도록 프로젝트 내부 전용 에셋 디렉토리에 배치합니다.")
    add_bullet_p(doc, " C:\\BeausCreators\\02.BSC_Branch\\projects\\Runnow\\assets\\careteam\\", "저장 위치:")
    add_bullet_p(doc, " 512x512 PNG (1:1 정사각형, 원형 마스크 최적화, 투명 배경 또는 딥 다크 네온 배경)", "이미지 규격:")
    add_bullet_p(doc, " 나이키 런 클럽 시그니처 형광 볼트(#CCFF00) + OLED 딥블랙(#08090C) 무드의 세련된 3D/일러스트 스타일", "디자인 무드:")
    add_bullet_p(doc, " leo_coach.png, luna_coach.png, ellie_diet.png, dr_kay.png (총 4종)", "파일명 표준:")

    output_filename = "2026-09-14_RUNNOW_AI_CareTeam_Agent_Profiles.docx"
    
    # 1) RunNow 프로젝트 폴더 저장
    proj_path = os.path.join(r"C:\BeausCreators\02.BSC_Branch\projects\Runnow", output_filename)
    doc.save(proj_path)
    print(f"[SUCCESS] Saved to Project: {proj_path}")

    # 2) 본사 보고서 폴더 저장 (SSOT)
    hq_dir = r"C:\BeausCreators\01.BSC_HQ\1.Documents\01.보고서\2026-09-14"
    os.makedirs(hq_dir, exist_ok=True)
    hq_path = os.path.join(hq_dir, output_filename)
    shutil.copyfile(proj_path, hq_path)
    print(f"[SUCCESS] Saved to HQ SSOT: {hq_path}")

    # 3) 바탕화면(Desktop) 복사본 저장
    desktop_dir = r"C:\Users\USER\Desktop"
    desktop_path = os.path.join(desktop_dir, output_filename)
    shutil.copyfile(proj_path, desktop_path)
    print(f"[SUCCESS] Saved to Desktop: {desktop_path}")

if __name__ == "__main__":
    main()
