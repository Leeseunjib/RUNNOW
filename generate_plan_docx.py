# -*- coding: utf-8 -*-
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Malgun Gothic'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Malgun Gothic'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Malgun Gothic'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r = p.add_run(text)
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Malgun Gothic'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r = p.add_run(text)
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_callout(doc, text, title="핵심 가이드"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"💡 {title}: ")
    r_title.font.name = 'Malgun Gothic'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

    r_body = p.add_run(text)
    r_body.font.name = 'Malgun Gothic'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def main():
    doc = docx.Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. 헤더 표지 블록
    tbl_header = doc.add_table(rows=1, cols=1)
    tbl_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_h = tbl_header.cell(0, 0)
    cell_h.width = Inches(6.8)
    set_cell_background(cell_h, "0F172A")
    set_cell_margins(cell_h, top=240, bottom=240, left=240, right=240)

    p_tag = cell_h.paragraphs[0]
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("BEAUSCREATORS TECHNICAL IMPLEMENTATION PLAN")
    r_tag.font.name = 'Malgun Gothic'
    r_tag.font.size = Pt(9)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

    p_t = cell_h.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(6)
    r_t = p_t.add_run("RunNow 1:1 전담 AI 케어팀 & 공공데이터 기반\n식단 관리 체크리스트 구축 상세 실행 플랜")
    r_t.font.name = 'Malgun Gothic'
    r_t.font.size = Pt(17)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_sub = cell_h.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    r_sub = p_sub.add_run("남/여 PT 음성 코칭 | 대화형 캘린더 자동 등록 | 식약처 Top 500+ 캐시 | 치팅 솔직 고백 만회 시스템")
    r_sub.font.name = 'Malgun Gothic'
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(8)
    p_meta.paragraph_format.space_after = Pt(12)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_m = p_meta.add_run("보고 대상: 이건우 대표님 (BeausCreators CEO) | 작성: BSC Core TF & RunNow 개발팀 | 일자: 2026.09.14")
    r_m.font.name = 'Malgun Gothic'
    r_m.font.size = Pt(8.5)
    r_m.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 1. 개요 및 사용자 검토 항목
    add_heading_1(doc, "1. 개요 및 대표님 사전 검토 항목")
    add_body_p(doc, "본 문서는 이건우 대표님께서 지시하신 '1:1 PT 대체/상생 포지셔닝', '공공기관(식약처 등) 영양 API 연동', '자책감 없는 식단 체크리스트 & 보정 관리 시스템'을 실제 앱에 구현하기 위한 세부 기술 작업 플랜입니다.")
    
    add_callout(doc, "기존 하단 5개 탭 [운동, 펫, 퀘스트, 상점, 설정] 중, 대표님의 핵심 킬러 피처인 [1:1 케어팀(🎯)]과 [식단 체크(🥗)]를 메인 탭으로 전진 배치하여 유저가 앱 실행 직후 1초 만에 케어를 체감할 수 있도록 구성합니다.", "UI 레이아웃 최적화 제안")

    # 2. 신규 생성 및 수정 파일 목록
    add_heading_1(doc, "2. 파일별 작업 명세 (Architecture Blueprint)")
    
    tbl_files = doc.add_table(rows=6, cols=3)
    tbl_files.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_files)

    headers = ["구분 / 파일명", "역할 및 주요 기능", "기술 스펙"]
    for i, h in enumerate(headers):
        cell = tbl_files.cell(0, i)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = 'Malgun Gothic'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    file_rows = [
        ("[신규] careTeam.js", 
         "• 3대 에이전트 인터랙션 (레오/루나 PT, 엘리 영양사, 닥터 케이)\n• 남/여 성별 및 음성(TTS) 토글\n• 대화형 주간 맞춤 운동 스케줄 자동 생성\n• 캘린더 타임라인 렌더링", 
         "Web Speech API (TTS), LocalStorage/Firestore 연동, 자연어 Intent 파서"),
        ("[신규] dietData.js", 
         "• 식품의약품안전처 통합영양성분 표준 DB 기반 Top 500+ 품목 인앱 번들\n• 외식/가공식품/자연식품 정밀 열량, 탄단지, 나트륨 데이터", 
         "0초 오프라인 캐시 (JSON), 온디바이스 초고속 검색"),
        ("[신규] dietManager.js", 
         "• 행동 기반 일일 식단 체크리스트 (아침/점심/간식/저녁)\n• [오늘 치팅/과식했어요] 솔직 고백 UX\n• 저녁 식단 자동 다운사이징 & 운동 보정 로직\n• 칼로리 수지 실시간 게이지", 
         "Dynamic Rescue 알고리즘, 칼로리 밸런스 미터, 다마고치 정직 보상 연동"),
        ("[수정] index.html", 
         "• 하단 탭 바 [케어팀] 및 [식단] 탭 추가\n• 1:1 코치실 및 일일 식단 체크리스트 뷰 섹션 마크업\n• 음식 검색 팝업 및 치팅 솔직 고백 모달 추가", 
         "HTML5 Semantic, 반응형 뷰 컨테이너"),
        ("[수정] styles.css & app.js", 
         "• Midnight Volt 테마 1:1 대화 말풍선, 캘린더 카드, 체크리스트 UI 스타일링\n• 탭 전환 라우팅 및 운동 완료 시 소모 칼로리 실시간 동기화 바인딩", 
         "CSS Grid/Flexbox, 모듈 초기화 이벤트 리스너")
    ]

    for row_idx, data in enumerate(file_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = tbl_files.cell(row_idx, col_idx)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Malgun Gothic'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    add_body_p(doc, "", space_after=6)

    # 3. 단계별 상세 작업 플랜 (Step-by-Step Execution Plan)
    add_heading_1(doc, "3. 단계별 세부 실행 플랜 (Step-by-Step Plan)")

    add_heading_2(doc, "Step 1. 식약처 표준 DB 인앱 캐시 및 식단 관리 모듈 구축 (dietData.js / dietManager.js)")
    add_bullet_p(doc, " 한국인이 가장 자주 찾는 다이어트/외식 500+개 식품 영양 데이터를 dietData.js로 번들링하여 검색 지연 0초 달성", "1.1 공공 DB 캐싱:")
    add_bullet_p(doc, " 아침(물+단백질20g), 점심(밥2/3+국물컷), 간식(당방어), 저녁(가벼운식사) 등 4대 행동 체크리스트 로직 구현", "1.2 일일 체크리스트:")
    add_bullet_p(doc, " 유저가 치팅을 클릭했을 때 자책감을 없애주는 격려 메시지와 함께, 저녁 식단을 자동으로 가볍게 보정하고 PT 캘린더에 15분 추가 운동을 제안하는 Rescue 알고리즘 구현", "1.3 치팅 솔직 고백 & 만회:")

    add_heading_2(doc, "Step 2. 1:1 케어팀 허브 및 대화형 스케줄러 개발 (careTeam.js)")
    add_bullet_p(doc, " 남성 코치 '레오'(파이팅), 여성 코치 '루나'(섬세함), 영양사 '엘리', 의사 '닥터 케이' 4인 프로필 및 대화 엔진 탑재", "2.1 에이전트 4인 페르소나:")
    add_bullet_p(doc, " '다음 주 월/수/금 저녁에 30분씩 운동할래' 입력 시 요일/시간/루틴을 자동 파싱하여 캘린더에 이벤트로 생성", "2.2 대화형 지능형 스케줄러:")
    add_bullet_p(doc, " 브라우저 내장 Web Speech API를 활용하여 추가 비용 0원으로 한국어 남/여 음성 코칭 출력", "2.3 오디오 TTS 탑재:")

    add_heading_2(doc, "Step 3. 메인 화면 UI 통합 및 스타일링 (index.html / styles.css / app.js)")
    add_bullet_p(doc, " 하단 탭 바에 케어팀과 식단 탭을 배치하고 기존 Midnight Volt 디자인 시스템 규격에 맞추어 네온 하이라이트 적용", "3.1 내비게이션 바 재배치:")
    add_bullet_p(doc, " 카카오톡처럼 친근한 1:1 챗 인터페이스와 주간 운동 타임라인 카드 렌더링", "3.2 1:1 코치실 뷰 구현:")
    add_bullet_p(doc, " 야식 방어 도장 클릭 시 tamagotchi.js와 연동하여 펫에게 숙면 버프 및 볼트 코인(VC) 지급", "3.3 다마고치 펫 연동:")

    add_heading_2(doc, "Step 4. 브라우저 실환경 전수 검증 및 워크스루 보고")
    add_bullet_p(doc, " 1:1 코치 대화 및 음성 출력, 캘린더 일정 생성, 식단 체크리스트 클릭, 치팅 고백 만회 플로우 정상 동작 검증", "4.1 E2E 기능 검증:")

    # 4. 검증 기준 및 인수 조건 (Acceptance Criteria)
    add_heading_1(doc, "4. 최종 인수 기준 (Acceptance Criteria)")
    add_bullet_p(doc, " [1:1 케어팀] 탭에서 남성 코치 레오와 여성 코치 루나를 원클릭으로 선택할 수 있어야 함", "AC-1 코치 선택:")
    add_bullet_p(doc, " '월수금 30분 플랜 짜줘' 대화 시 캘린더에 해당 요일 루틴 카드가 즉시 생성되어야 함", "AC-2 대화형 캘린더:")
    add_bullet_p(doc, " 식단 검색창에 '삼겹살', '서브웨이', '닭가슴살' 입력 시 식약처 공인 칼로리/탄단지가 즉각 노출되어야 함", "AC-3 식약처 DB 연동:")
    add_bullet_p(doc, " [오늘 치팅/과식했어요] 클릭 시 자책감 완화 메시지와 함께 저녁 체크리스트 및 운동 보정이 연동되어야 함", "AC-4 치팅 만회 시스템:")

    add_body_p(doc, "", space_after=8)
    add_callout(doc, "이건우 대표님께서 본 구현 플랜을 검토 후 승인(Proceed)해 주시면, 즉시 1단계 모듈 구현부터 한 치의 오차도 없이 신속히 완수하겠습니다.", "대표님 승인 대기")

    output_filename = "2026-09-14_RUNNOW_CareTeam_Implementation_Plan.docx"
    doc.save(output_filename)
    print(f"[SUCCESS] Created {output_filename}")

if __name__ == "__main__":
    main()
